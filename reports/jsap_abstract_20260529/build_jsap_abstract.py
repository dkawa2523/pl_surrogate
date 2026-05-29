from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt


OUT = Path(__file__).resolve().parent
DOCX_PATH = OUT / "jsap_icp_structure_feature_abstract.docx"
MD_PATH = OUT / "jsap_icp_structure_feature_abstract_source.md"
FIG_PATH = OUT / "jsap_chamber_prediction_figure.png"
SUMMARY_PATH = OUT / "jsap_abstract_build_summary.json"

TITLE_JP = "構造特徴量に基づく低圧ICPコイル設計のサロゲート解析"
TITLE_EN = "Structure-feature-based surrogate analysis for coil design in low-pressure ICP"
AUTHORS_JP = "○発表者氏名^1，共著者氏名^1,2（投稿前に差し替え）"
AUTHORS_EN = "○Presenter Name^1, Co-author Name^1,2 (to be replaced before submission)"
AFFIL_JP = "^1 所属1，^2 所属2（投稿前に差し替え）"
AFFIL_EN = "^1 Affiliation 1, ^2 Affiliation 2 (to be replaced before submission)"
EMAIL = "E-mail: presenter@example.com（投稿前に差し替え）"

BODY_PARAGRAPHS = (
    "半導体プラズマプロセスに用いられる誘導結合プラズマ（ICP）では，RFコイルの形状や配置が"
    "誘導電磁場，電子加熱，電離分布を通じてプラズマ密度の空間分布を決める[1]。したがってコイル設計では，"
    "単一の代表値だけでなく，チャンバー内に形成される二次元場そのものを予測し，ウェハ近傍での均一性や"
    "密度維持を同時に評価する必要がある。従来の機械学習では，コイル幅，位置，本数などを数値パラメータとして"
    "入力することが多い。しかし，コイル本数や配置の違いは単なる数値の変化ではなく，電磁場とプラズマが解かれる"
    "幾何構造の変化である。この違いを無視すると，物理的には近い構造が別の入力として扱われたり，未知の配置に対する"
    "外挿を過信したりする問題が生じる。",
    "本研究では，GEC reference cell 型ICPを基にした軸対称二次元数値データを用い[2,3]，コイルを寸法列としてではなく，"
    "チャンバー内の空間構造特徴として表現するサロゲート解析を行う。すなわち，コイル配置から境界や近接度を表す"
    "構造場を作り，運転条件と組み合わせて電子密度，イオン密度，電子温度，電位などの二次元プラズマ場を予測する。"
    "モデルにはニューラルオペレータ型のサロゲート[4]を用い，有限次元の設計変数から値を直接回帰するのではなく，"
    "構造場からプラズマ場への写像として学習する。密度の桁差や物理量ごとのスケール差は，学習データのみから決めた"
    "前処理で扱う。",
    "予備的な学習評価では，構造特徴量を用いることで，コイル寸法を単なるパラメータとして扱う場合に比べて，"
    "密度ピークや勾配を含む空間分布を安定に再現できる傾向が得られた。特に，密度や温度のようにプラズマ領域で"
    "意味を持つ量と，磁場のようにコイル近傍まで広がる量では，評価すべき領域が異なる。この点を明示して学習・評価することも，"
    "サロゲートを物理的に解釈するうえで重要である。",
    "さらに，学習済みサロゲートを用いて，電子密度の均一性を高めつつ密度低下を避けるコイル候補の探索を行う。"
    "ここでも最適化する対象は，寸法値そのものではなく，候補形状から作られる構造特徴を通じて評価されるプラズマ場である。"
    "これにより，数値上の最良点だけでなく，物理的に異なる複数の候補を比較できる。発表では，場分布の予測例と候補探索の結果を示し，"
    "サロゲートを最終解ではなく，物理制約を満たす再計算候補を効率よく絞り込む道具として位置づける。",
)

REFS = (
    "[1] M. A. Lieberman and A. J. Lichtenberg, Principles of Plasma Discharges and "
    "Materials Processing, Wiley (2005).  [2] P. J. Hargis Jr. et al., Rev. Sci. "
    "Instrum. 65, 140-154 (1994), doi:10.1063/1.1144770.  [3] P. A. Miller et al., "
    "J. Res. Natl. Inst. Stand. Technol. 100, 427-439 (1995), doi:10.6028/jres.100.032.  "
    "[4] Z. Li et al., Fourier Neural Operator "
    "for Parametric PDEs, ICLR (2021)."
)


def _set_run_font(run, size_pt: float = 9.0, bold: bool = False) -> None:
    run.font.name = "Times New Roman"
    run.font.size = Pt(size_pt)
    run.font.bold = bold
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), "ＭＳ 明朝")
    rfonts.set(qn("w:ascii"), "Times New Roman")
    rfonts.set(qn("w:hAnsi"), "Times New Roman")


def _add_para(
    doc: Document,
    text: str = "",
    *,
    size: float = 9.0,
    bold: bool = False,
    align: WD_ALIGN_PARAGRAPH | None = None,
    space_after: float = 0,
    first_indent: bool = False,
):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.0
    if first_indent:
        p.paragraph_format.first_line_indent = Mm(3.5)
    run = p.add_run(text)
    _set_run_font(run, size, bold)
    return p


def write_markdown() -> None:
    lines = [
        f"# {TITLE_JP}",
        "",
        TITLE_EN,
        "",
        AUTHORS_JP,
        "",
        AFFIL_JP,
        "",
        "",
    ]
    for body in BODY_PARAGRAPHS:
        lines.extend([body, ""])
    lines.extend(
        [
            f"![チャンバー構造とサロゲートが予測した電子密度分布]({FIG_PATH.name})",
            "",
            "**図1** チャンバー構造，コイル配置，およびサロゲートが予測した電子密度分布の例。"
            "赤枠はRFコイル，破線は均一性評価に用いた代表断面を示す。",
            "",
            "**参考文献**",
            REFS,
            "",
            "投稿前差し替え: 発表者氏名，共著者氏名，所属，E-mail，講演番号。",
        ]
    )
    MD_PATH.write_text("\n".join(lines), encoding="utf-8")


def write_docx() -> None:
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.top_margin = Mm(14)
    sec.bottom_margin = Mm(14)
    sec.left_margin = Mm(15)
    sec.right_margin = Mm(15)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(9)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "ＭＳ 明朝")

    _add_para(doc, TITLE_JP, size=12, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=1)
    _add_para(doc, TITLE_EN, size=9.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    _add_para(doc, AUTHORS_JP, size=9, align=WD_ALIGN_PARAGRAPH.CENTER)
    _add_para(doc, AUTHORS_EN, size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    _add_para(doc, AFFIL_JP, size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    _add_para(doc, AFFIL_EN, size=8, align=WD_ALIGN_PARAGRAPH.CENTER)
    _add_para(doc, EMAIL, size=8, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)

    body_sec = doc.add_section(WD_SECTION_START.CONTINUOUS)
    sect_pr = body_sec._sectPr
    cols = sect_pr.xpath("./w:cols")
    cols = cols[0] if cols else OxmlElement("w:cols")
    if not sect_pr.xpath("./w:cols"):
        sect_pr.append(cols)
    cols.set(qn("w:num"), "2")
    cols.set(qn("w:space"), "420")

    for body in BODY_PARAGRAPHS[:2]:
        _add_para(doc, body, size=8.2, first_indent=True, space_after=2)

    doc.add_picture(str(FIG_PATH), width=Inches(3.12))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.paragraphs[-1].paragraph_format.space_after = Pt(2)
    _add_para(
        doc,
        "図1　チャンバー構造，コイル配置，およびサロゲートが予測した電子密度分布の例。",
        size=7.2,
        bold=True,
        space_after=2,
    )

    for body in BODY_PARAGRAPHS[2:]:
        _add_para(doc, body, size=8.2, first_indent=True, space_after=2)

    _add_para(doc, "参考文献", size=8.0, bold=True)
    _add_para(doc, REFS, size=7.0)
    _add_para(doc, "投稿前差し替え: 発表者氏名，共著者氏名，所属，E-mail，講演番号。", size=7.0)

    doc.core_properties.title = TITLE_JP
    doc.core_properties.subject = "応用物理学会予稿案"
    doc.core_properties.keywords = (
        "ICP, neural operator, surrogate model, structure features, coil design"
    )
    doc.core_properties.comments = (
        "Draft generated from project report artifacts; author and affiliation placeholders "
        "must be replaced before submission."
    )
    doc.save(DOCX_PATH)


def write_summary() -> None:
    summary = {
        "docx": str(DOCX_PATH),
        "source_markdown": str(MD_PATH),
        "figure": str(FIG_PATH),
        "official_template_downloaded": str(OUT / "jsap_official_template.doc"),
        "format_basis": {
            "JSAP_template_page": "https://meeting.jsap.or.jp/template/index.html",
            "notes": (
                "Official JSAP template page states 1 page and no character-count limit. "
                "This draft uses A4, compact heading block, and two-column body."
            ),
        },
        "needs_replacement_before_submission": [
            "presenter name",
            "coauthor names",
            "affiliations",
            "email",
            "presentation number if assigned",
        ],
    }
    SUMMARY_PATH.write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")


if __name__ == "__main__":
    write_markdown()
    write_docx()
    write_summary()
    print(json.dumps(json.loads(SUMMARY_PATH.read_text(encoding="utf-8")), ensure_ascii=False, indent=2))
